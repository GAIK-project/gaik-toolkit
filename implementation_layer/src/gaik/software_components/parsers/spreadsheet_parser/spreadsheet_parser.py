"""Spreadsheet (.xlsx, .csv) to Markdown tables that keep the source row numbers."""

from __future__ import annotations

import csv
from collections.abc import Iterable
from datetime import datetime, time
from pathlib import Path
from typing import Any


def _cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        text = value.date().isoformat() if value.time() == time() else value.isoformat(sep=" ")
    elif isinstance(value, float) and value.is_integer():
        text = str(int(value))
    else:
        text = str(value)
    return " ".join(text.splitlines()).replace("|", "\\|")


def _table(rows: Iterable[Iterable[Any]]) -> str | None:
    """Render rows as a Markdown table with a leading ``Row`` column; None if all are empty."""
    kept = []
    for number, row in enumerate(rows, start=1):
        cells = [_cell(v) for v in row]
        while cells and not cells[-1].strip():
            cells.pop()
        if cells:
            kept.append([str(number), *cells])
    if not kept:
        return None
    width = max(len(cells) for cells in kept)
    header, *body = [cells + [""] * (width - len(cells)) for cells in kept]
    lines = [["Row", *header[1:]], ["---"] * width, *body]
    return "\n".join("| " + " | ".join(line) + " |" for line in lines)


def _xlsx_blocks(path: Path) -> list[str]:
    """One ``## Sheet`` table per non-empty sheet, with the stored results of formulas."""
    from openpyxl import load_workbook

    values = load_workbook(path, read_only=True, data_only=True)
    formulas = load_workbook(path, read_only=True)
    try:
        blocks, results = [], []
        for sheet, formula_sheet in zip(values.worksheets, formulas.worksheets):
            rows = list(sheet.iter_rows(values_only=True))
            for row, cells in zip(rows, formula_sheet.iter_rows()):
                results += [v for v, c in zip(row, cells) if c.data_type == "f"]
            if table := _table(rows):
                blocks.append(f"## Sheet: {sheet.title}\n\n{table}")
    finally:
        values.close()
        formulas.close()
    # A file written by a library such as openpyxl has formulas but no results, which
    # would read as blank cells. A single empty result can be a legitimate "".
    if results and all(v is None for v in results):
        raise ValueError(
            f"{path.name}: its formulas have no stored results; "
            "open and save it in Excel or LibreOffice to calculate them"
        )
    return blocks


class SpreadsheetParser:
    """Convert an .xlsx workbook or a .csv file to Markdown tables.

    The first non-empty row of each sheet is the header. A ``Row`` column holds the
    1-based sheet or CSV row number of each data row, so a citation can name the row.
    """

    def parse_spreadsheet(self, file_path: str | Path) -> str:
        """Return the spreadsheet as Markdown: one ``## Sheet: <name>`` table per sheet."""
        path = Path(file_path)
        suffix = path.suffix.lower()
        if suffix == ".xlsx":
            blocks = _xlsx_blocks(path)
        elif suffix == ".csv":
            with path.open(encoding="utf-8-sig", newline="") as handle:
                blocks = [table] if (table := _table(csv.reader(handle))) else []
        else:
            raise ValueError(f"{path.name}: unsupported spreadsheet type {suffix}")
        if not blocks:
            raise ValueError(f"{path.name}: spreadsheet has no non-empty rows")
        return "\n\n".join(blocks)

    def parse_document(self, file_path: str | Path) -> dict[str, Any]:
        """Parse a spreadsheet and return its Markdown with file metadata."""
        path = Path(file_path)
        text_content = self.parse_spreadsheet(path)
        return {
            "file_path": str(path),
            "file_name": path.name,
            "file_extension": path.suffix.lower(),
            "text_content": text_content,
            "content_length": len(text_content),
            "word_count": len(text_content.split()),
            "parsing_method": "spreadsheet",
        }
