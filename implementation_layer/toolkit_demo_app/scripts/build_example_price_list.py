"""Write the example price list as JSON, for the Purchase Order demo's preview.

The browser cannot read a spreadsheet, so the preview of "price list.xlsx" in the example
section shows a JSON copy of its first sheet:

    public/data/po-bom-example/price-list.json

Run it after changing the price list:

    uv run python implementation_layer/toolkit_demo_app/scripts/build_example_price_list.py

`--check` exits non-zero when the JSON no longer matches the spreadsheet; a unit test runs it.
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

from openpyxl import load_workbook

APP_DIR = Path(__file__).resolve().parent.parent
SOURCE = APP_DIR / "public" / "data" / "po-bom-example" / "price list.xlsx"
TARGET = APP_DIR / "public" / "data" / "po-bom-example" / "price-list.json"


def read_rows(path: Path = SOURCE) -> list[list[str | int | float | None]]:
    """The rows of the first sheet, without the empty cells at the end of each row."""
    workbook = load_workbook(path, data_only=True)
    try:
        sheet = workbook[workbook.sheetnames[0]]
        rows: list[list[str | int | float | None]] = []
        for values in sheet.iter_rows(values_only=True):
            row = [v if not isinstance(v, str) else v.strip() or None for v in values]
            while row and row[-1] is None:
                row.pop()
            rows.append(row)
        return rows
    finally:
        workbook.close()


def render() -> str:
    return json.dumps({"rows": read_rows()}, ensure_ascii=False, indent=1) + "\n"


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--check", action="store_true", help="fail when the JSON is stale")
    args = parser.parse_args()

    expected = render()
    if args.check:
        current = TARGET.read_text(encoding="utf-8") if TARGET.exists() else ""
        if current.replace("\r\n", "\n") != expected:
            print(f"{TARGET} is stale; run this script without --check.", file=sys.stderr)
            return 1
        return 0
    TARGET.write_text(expected, encoding="utf-8", newline="\n")
    print(f"Wrote {TARGET}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
