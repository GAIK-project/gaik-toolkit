"""Generate the synthetic house condition assessment dataset.

Writes the customer documents to ``inputs/`` and ``sample_report.docx`` next to it, then
turns the recording scripts into mp3 files with TextToSpeech. The audio step is an API
call and needs credentials in the environment or a ``.env`` file; ``--documents-only``
skips it. All content is fixed and fictional.
"""

from __future__ import annotations

import argparse
import logging
import time
from datetime import date
from pathlib import Path

import openpyxl
import pymupdf
from docx import Document
from dotenv import load_dotenv
from gaik.software_components.text_to_speech import TextToSpeech
from PIL import Image, ImageDraw, ImageFont

HERE = Path(__file__).resolve().parent
INPUTS = HERE / "inputs"
RECORDINGS = ("rec1_exterior_attic", "rec2_interior_wet_rooms", "rec3_building_services")
ADDRESS = "12 Example Road, Sampleton"

# The TTS endpoint occasionally drops a connection; retried with backoff, not silently
# skipped, so a real failure still stops the script.
MAX_RETRIES = 4
BACKOFF_BASE_SECONDS = 2.0

log = logging.getLogger("generate_dataset")


def write_pdf(path: Path, title: str, pages: list[list[tuple[str, str | list]]]) -> None:
    """Write A4 pages of ``(style, content)`` blocks: ``h1``, ``h2``, ``p`` or ``table``.

    Uses the built-in Helvetica, which is not shaped: HTML rendering would put ligatures
    such as "ﬂ" in the text layer and break the verbatim quote checks downstream.
    """
    regular, bold = pymupdf.Font("helv"), pymupdf.Font("hebo")
    styles = {"h1": (bold, 17), "h2": (bold, 13), "p": (regular, 11)}
    doc = pymupdf.open()
    for number, blocks in enumerate(pages, start=1):
        page = doc.new_page()
        writer = pymupdf.TextWriter(page.rect)
        writer.append((60, 45), f"{title} | {ADDRESS}", font=regular, fontsize=9)
        writer.append((60, 810), f"Page {number} of {len(pages)}", font=regular, fontsize=9)
        y = 70
        for style, content in blocks:
            if style == "table":
                second = 80 + max(bold.text_length(row[0], fontsize=11) for row in content)
                for i, row in enumerate(content):
                    y += 18
                    for x, cell in zip((60, second), row, strict=True):
                        writer.append((x, y), cell, font=bold if i == 0 else regular, fontsize=11)
            else:
                font, size = styles[style]
                box = pymupdf.Rect(60, y, 535, 780)
                writer.fill_textbox(box, content, font=font, fontsize=size, warn=False)
                y = writer.last_point.y
            y += 12
            if y > 780:
                raise RuntimeError(f"Page {number} of {path.name} does not fit on the page")
        writer.write_text(page)
    doc.save(path, no_new_id=True)


def write_renovation_report_1998(path: Path) -> None:
    page1 = [
        ("h1", "Renovation Report 1998"),
        (
            "p",
            f"Property: {ADDRESS}\nPrepared by: Sampleton Building Services Ltd\n"
            "Date: 30 October 1998",
        ),
        ("h2", "1. Property"),
        (
            "p",
            "The property is a 1½-storey timber-frame detached house built in 1978. A partial "
            "basement houses the technical room.",
        ),
        ("h2", "2. Original structures (1978)"),
        (
            "table",
            [
                ("Part", "Description"),
                ("Foundation", "Cast concrete strip foundations and plinth"),
                ("Base floor", "Concrete slab on ground"),
                ("External walls", "Timber frame, 100 mm mineral wool, painted board cladding"),
                ("Roof", "Concrete tile roof with an underlay, on timber trusses"),
            ],
        ),
        (
            "p",
            "The original concrete tile roof was in fair condition and was not part of the 1998 "
            "renovation.",
        ),
    ]
    page2 = [
        ("h2", "3. Scope of the renovation"),
        (
            "p",
            "The renovation was carried out between 1 June and 30 October 1998. It covered:\n"
            "• waterproofing and tiling of the bathroom and the sauna floor;\n"
            "• a new mechanical ventilation system;\n"
            "• replacement of the main distribution board.",
        ),
        ("h2", "4. Bathroom and sauna"),
        (
            "p",
            "The old tiles in the bathroom were removed down to the concrete. The bathroom floor "
            "and walls and the sauna floor were waterproofed in 1998 with a brush-applied "
            "waterproofing membrane, which was lapped into new plastic floor drains with clamping "
            "rings. New floor and wall tiles were laid on the membrane.",
        ),
        ("p", "The utility room and the WC were not renovated."),
    ]
    page3 = [
        ("h2", "5. Ventilation"),
        (
            "p",
            "The original natural ventilation was replaced with mechanical supply and extract "
            "ventilation. Mechanical ventilation system installed during 1998 renovation. Design "
            "life 25 years.",
        ),
        (
            "p",
            "The supply-air unit is in the basement technical room. It supplies filtered outdoor "
            "air to the living room and the bedrooms. Air is extracted through ducts from the "
            "kitchen, the bathroom, the sauna, the WC and the utility room. The filters should be "
            "changed at least once a year.",
        ),
        ("h2", "6. Electrical"),
        (
            "p",
            "The main distribution board was replaced in 1998. The new board has automatic "
            "circuit breakers and a residual current device for the wet room circuits. The rest "
            "of the wiring is original from 1978.",
        ),
    ]
    page4 = [
        ("h2", "7. Costs"),
        (
            "table",
            [
                ("Work", "Cost (FIM, incl. VAT)"),
                ("Bathroom and sauna waterproofing and tiling", "36 500"),
                ("Mechanical ventilation system", "42 800"),
                ("Main distribution board replacement", "14 200"),
                ("Design and site supervision", "9 600"),
                ("Total", "103 100"),
            ],
        ),
        ("h2", "8. Handover"),
        (
            "p",
            "The work was completed and handed over to the owner on 30 October 1998.\n"
            "Site supervisor, Sampleton Building Services Ltd",
        ),
    ]
    write_pdf(path, "Renovation Report 1998", [page1, page2, page3, page4])


def write_roof_renovation_2012(path: Path) -> None:
    page = [
        ("h1", "Roof Renovation 2012"),
        (
            "p",
            f"Property: {ADDRESS}\nContractor: Sampleton Roofing Ltd\nCompleted: 21 September 2012",
        ),
        ("h2", "1. Work done"),
        (
            "p",
            "• The original concrete tiles and battens were removed.\n"
            "• The old underlay was removed and a new roof underlay was installed over the whole "
            "roof.\n"
            "• New counter-battens and battens were fitted.\n"
            "• The new roof covering is profiled steel sheet, 0.5 mm, coated, dark grey.\n"
            "• New eaves gutters, downpipes and snow guards were installed.",
        ),
        ("h2", "2. Trusses"),
        (
            "p",
            "The trusses were checked while the roof was open. They were sound and needed no "
            "repairs.",
        ),
    ]
    write_pdf(path, "Roof Renovation 2012", [page])


def write_moisture_survey_2019(path: Path) -> None:
    doc = Document()
    doc.add_heading("Moisture Survey of Wet Rooms 2019", level=1)
    doc.add_paragraph(
        f"Property: {ADDRESS}\nSurvey date: 6 May 2019\nSurveyor: Sampleton Moisture Surveys Ltd"
    )
    doc.add_heading("Method", level=2)
    doc.add_paragraph(
        "Surface moisture meter, relative scale 0–200. Readings below 60 are normal, 60 to 90 "
        "slightly elevated and above 90 elevated. The reference reading on the dry hallway "
        "floor was 28."
    )
    doc.add_heading("Readings", level=2)
    table = doc.add_table(rows=0, cols=3, style="Table Grid")
    for row in [
        ("Location", "Reading", "Assessment"),
        ("Bathroom floor, next to the shower floor drain", "34", "Normal"),
        ("Bathroom floor, middle", "31", "Normal"),
        ("Bathroom, shower wall", "30", "Normal"),
        ("Sauna floor", "33", "Normal"),
        ("Utility room floor, next to the floor drain", "36", "Normal"),
        ("WC floor", "29", "Normal"),
    ]:
        for cell, text in zip(table.add_row().cells, row, strict=True):
            cell.text = text
    doc.add_heading("Conclusion", level=2)
    doc.add_paragraph(
        "No elevated readings were found in the wet rooms. The waterproofing was working at the "
        "time of the survey."
    )
    doc.save(path)


def write_maintenance_log(path: Path) -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Log"
    sheet.append(["Date", "Entry"])
    for row in [
        (date(2014, 5, 10), "Hedge trimmed along the street side"),
        (date(2014, 10, 4), "Ventilation filters changed"),
        (date(2015, 6, 13), "Lawn fertilised and aerated"),
        (date(2015, 10, 10), "Ventilation filters changed"),
        (
            date(2016, 4, 22),
            "Sewer flushing: main sewer line flushed from the basement cleanout to the street "
            "by a drain cleaning company",
        ),
        (date(2016, 8, 20), "Hedge trimmed"),
        (date(2017, 5, 6), "Apple tree pruned and flower beds mulched"),
        (date(2017, 10, 14), "Ventilation filters changed"),
        (date(2019, 10, 12), "Ventilation filters changed"),
        (date(2020, 6, 6), "Lawn mower serviced before the mowing season"),
        (date(2021, 10, 9), "Ventilation filters changed"),
        (date(2022, 7, 2), "Hedge trimmed"),
        (date(2023, 10, 21), "Ventilation filters changed"),
        (date(2024, 5, 18), "Bare patches in the lawn reseeded"),
        (date(2025, 10, 18), "Ventilation filters changed"),
    ]:
        sheet.append(row)
        sheet.cell(sheet.max_row, 1).number_format = "yyyy-mm-dd"
    sheet.column_dimensions["A"].width = 12
    sheet.column_dimensions["B"].width = 95
    workbook.save(path)


def write_floor_plan(path: Path) -> None:
    floors = {
        "Ground floor": [
            ("Living room", 30.0, (40, 140, 460, 440)),
            ("Kitchen", 14.0, (460, 140, 760, 320)),
            ("Bedroom 1", 12.0, (460, 320, 760, 440)),
            ("Entrance hall", 7.0, (40, 440, 220, 680)),
            ("WC", 2.0, (220, 440, 330, 680)),
            ("Bathroom", 6.0, (330, 440, 490, 680)),
            ("Sauna", 4.0, (490, 440, 620, 680)),
            ("Utility room", 6.5, (620, 440, 760, 680)),
        ],
        "Upper floor": [
            ("Bedroom 2", 14.0, (840, 140, 1160, 360)),
            ("Bedroom 3", 13.0, (1160, 140, 1460, 360)),
            ("Study", 10.0, (840, 360, 1080, 580)),
            ("Upper hall", 9.0, (1080, 360, 1310, 580)),
            ("Storage", 2.5, (1310, 360, 1460, 580)),
        ],
        "Partial basement": [("Technical room", 12.0, (840, 660, 1140, 800))],
    }
    image = Image.new("RGB", (1500, 840), "white")
    draw = ImageDraw.Draw(image)
    title, label, font = (ImageFont.load_default(size) for size in (34, 24, 20))
    draw.text((40, 30), f"Floor plan: {ADDRESS}", fill="black", font=title)
    for floor, rooms in floors.items():
        left = min(box[0] for _, _, box in rooms)
        top = min(box[1] for _, _, box in rooms)
        draw.text((left, top - 35), floor, fill="black", font=label)
        for name, area, box in rooms:
            draw.rectangle(box, outline="black", width=4)
            centre = ((box[0] + box[2]) / 2, (box[1] + box[3]) / 2)
            text = f"{name}\n{area:.1f} m2"
            draw.multiline_text(centre, text, fill="black", font=font, anchor="mm", align="center")
    living = sum(area for floor in ("Ground floor", "Upper floor") for _, area, _ in floors[floor])
    draw.text((40, 730), f"Living area, ground and upper floor: {living:.1f} m2", "black", label)
    image.save(path)


def write_sample_report(path: Path) -> None:
    """A report on a different, fictional property: for tone and format only."""
    field = "[field observation, 3 February 2026]"
    doc = Document()
    doc.add_heading("Condition assessment – 8 Specimen Street, Modelby", level=1)
    doc.add_paragraph("Inspection: 3 February 2026. Inspector: Alex Sample.")
    sections = {
        "Summary": [
            "The house was in fair condition for its age. The brick walls and the roof covering "
            "were in good condition. The bathroom waterproofing dates from 1983 and is past its "
            "expected service life. The oil boiler worked, but the underground oil tank has not "
            "been inspected. The attic was not inspected because the hatch was screwed shut. The "
            "most urgent actions are to renew the bathroom waterproofing and to have the oil tank "
            "inspected."
        ],
        "Property and background": [
            "According to the 1983 renovation invoice, the property is a single-storey brick "
            "house built in 1965 [1983 renovation invoice]. Floor area: (missing: floor area).",
            "According to the 1983 renovation invoice, the bathroom was renovated in 1983 [1983 "
            "renovation invoice]. According to the 2015 roofing invoice, the roof covering was "
            "renewed in 2015 [2015 roofing invoice].",
        ],
        "Structures and roof": [
            f"It was observed that the brick walls had no cracks or loose bricks {field}. The "
            f"concrete plinth showed minor spalling at the north-east corner {field}.",
            "According to the 2015 roofing invoice, the roof covering is bitumen felt [2015 "
            "roofing invoice]. The attic and the roof underlay were not inspected because the "
            "attic hatch was screwed shut.",
        ],
        "Wet rooms": [
            "According to the 1983 renovation invoice, the bathroom waterproofing dates from 1983 "
            "[1983 renovation invoice]. Moisture measurements made during the visit showed normal "
            f"readings on the bathroom floor and walls {field}."
        ],
        "Building services (heating, ventilation, water and sewer)": [
            "It was observed that the oil boiler in the boiler room ran normally, with soot "
            f"around the burner hatch {field}. According to the 2004 installation record, the "
            "boiler was installed in 2004 [2004 installation record]. No inspection record was "
            "available for the 3,000-litre underground oil tank.",
            f"It was observed that the house has natural ventilation through wall vents {field}. "
            "According to the 2021 sewer camera inspection report, the sewer pipes were in fair "
            "condition, with minor settlement near the house wall [2021 sewer camera inspection "
            "report].",
        ],
        "Electrical systems": [
            "It was observed that the main distribution board had screw fuses and no residual "
            f"current device {field}. According to the 2019 electrical inspection record, no "
            "faults requiring immediate repair were found [2019 electrical inspection record]."
        ],
    }
    for title, paragraphs in sections.items():
        doc.add_heading(title, level=2)
        for text in paragraphs:
            doc.add_paragraph(text)
    doc.add_heading("Recommendations", level=2)
    for text in [
        "Repair now: Clean the soot from the boiler room and have the burner adjusted.",
        "Within 1–2 years: Renew the bathroom waterproofing.",
        "Within 3–5 years: Add a residual current device to the main distribution board.",
        "Further investigation: Have the underground oil tank inspected.",
        "Further investigation: Open the attic hatch and inspect the attic and the underlay.",
        "Maintenance: Have the oil burner serviced every year.",
    ]:
        doc.add_paragraph(text, style="List Bullet")
    doc.save(path)


def _is_permanent(exc: Exception) -> bool:
    """A config/input problem, or a 4xx response: retrying can't fix either."""
    if isinstance(exc, ValueError):
        return True
    status = getattr(getattr(exc, "response", None), "status_code", None)
    return status is not None and 400 <= status < 500


def write_recordings(use_azure: bool) -> None:
    """Synthesize each recording script to mp3. An API call that needs credentials.

    Retries a failed call with exponential backoff, since the TTS endpoint occasionally
    drops a connection; a config error or a 4xx response raises immediately instead.
    """
    load_dotenv()
    # "fable" is a classic OpenAI voice that the "tts-hd" deployment accepts; gaik's own
    # default voice ("marin") is for newer models, and Azure rejects it on this deployment.
    tts = TextToSpeech(use_azure=use_azure, language="en", voice="fable")
    for name in RECORDINGS:
        text = (HERE / "recording_scripts" / f"{name}.txt").read_text(encoding="utf-8")
        for attempt in range(1, MAX_RETRIES + 1):
            try:
                result = tts.synthesize(text, language="en")
                break
            except Exception as exc:
                if _is_permanent(exc) or attempt == MAX_RETRIES:
                    raise
                delay = BACKOFF_BASE_SECONDS**attempt
                log.warning(
                    "%s: TTS attempt %d/%d failed, retrying in %.0fs",
                    name,
                    attempt,
                    MAX_RETRIES,
                    delay,
                    exc_info=True,
                )
                time.sleep(delay)
        print(f"Wrote {result.save(INPUTS / f'{name}.mp3')}")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--documents-only", action="store_true", help="Skip the text-to-speech (API) step"
    )
    parser.add_argument(
        "--openai", action="store_true", help="Use OpenAI instead of Azure OpenAI for speech"
    )
    args = parser.parse_args()

    INPUTS.mkdir(exist_ok=True)
    write_renovation_report_1998(INPUTS / "renovation_report_1998.pdf")
    write_roof_renovation_2012(INPUTS / "roof_renovation_2012.pdf")
    write_moisture_survey_2019(INPUTS / "moisture_survey_2019.docx")
    write_maintenance_log(INPUTS / "maintenance_log.xlsx")
    write_floor_plan(INPUTS / "floor_plan.png")
    write_sample_report(HERE / "sample_report.docx")
    print(f"Wrote the documents to {INPUTS} and {HERE / 'sample_report.docx'}")
    if not args.documents_only:
        write_recordings(use_azure=not args.openai)


if __name__ == "__main__":
    main()
