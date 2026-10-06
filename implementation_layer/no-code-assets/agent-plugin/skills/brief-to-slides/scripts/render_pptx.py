"""Render a deck spec to an editable .pptx inside a COPY of a template.

- Removes every slide already in the template (and orphan slide relationships), so no
  old text, notes or hidden slides survive.
- Uses the template's layout (THEME['layout']) and its TITLE placeholder for slide titles.
- Writes speaker notes into the PowerPoint Notes pane (creates the notes placeholder
  when the template's notes master has none).
"""
import copy
import os
import zipfile
from lxml import etree
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from dsl import *

ICON_DIR = os.environ.get('LUCIDE_DIR', os.path.expanduser('~/.cache/brief-to-slides/node_modules/lucide-static/icons'))
CACHE = os.path.expanduser('~/.cache/brief-to-slides/iconpng')
ALIGN = {'l': PP_ALIGN.LEFT, 'c': PP_ALIGN.CENTER, 'r': PP_ALIGN.RIGHT}
VAL = {'t': MSO_ANCHOR.TOP, 'm': MSO_ANCHOR.MIDDLE, 'b': MSO_ANCHOR.BOTTOM}
NS = ('xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
      'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"')
NOTES_BODY = (f'<p:sp {NS}><p:nvSpPr><p:cNvPr id="3" name="Notes Placeholder 2"/><p:cNvSpPr><a:spLocks noGrp="1"/></p:cNvSpPr>'
              '<p:nvPr><p:ph type="body" idx="1"/></p:nvPr></p:nvSpPr><p:spPr/><p:txBody><a:bodyPr/><a:lstStyle/><a:p/></p:txBody></p:sp>')


def icon_svg(name, color, sw=2.0):
    path = f'{ICON_DIR}/{name}.svg'
    if not os.path.exists(path):
        raise FileNotFoundError(f'Lucide icon "{name}" not found in {ICON_DIR} (check the name on lucide.dev or run setup_env.sh)')
    s = open(path).read()
    s = s[s.index('<svg'):]
    return s.replace('stroke="currentColor"', f'stroke="#{color}"').replace('stroke-width="2"', f'stroke-width="{sw}"')


def icon_png(name, color, sw=2.0):
    import cairosvg
    os.makedirs(CACHE, exist_ok=True)
    p = f'{CACHE}/{name}-{color}-{sw}.png'
    if not os.path.exists(p):
        cairosvg.svg2png(bytestring=icon_svg(name, color, sw).encode(), write_to=p, output_width=256, output_height=256)
    return p


def rgb(h):
    return RGBColor.from_string(h)


def strip_style(shape):
    """Remove the theme style reference so shapes get no theme shadow/outline."""
    st = shape._element.find(qn('p:style'))
    if st is not None:
        shape._element.remove(st)


def fill_tf(tf, el):
    tf.word_wrap = True
    bp = tf._txBody.find(qn('a:bodyPr'))
    for k in ('lIns', 'tIns', 'rIns', 'bIns'):
        bp.set(k, '0')
    tf.vertical_anchor = VAL[el['valign']]
    for child in list(bp):
        if child.tag in (qn('a:spAutoFit'), qn('a:normAutofit'), qn('a:noAutofit')):
            bp.remove(child)
    etree.SubElement(bp, qn('a:noAutofit'))
    first = True
    for para in el['paras']:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        p.alignment = ALIGN[para.get('align', el['align'])]
        p.line_spacing = para.get('lh', el['lh'])
        if not first:
            p.space_before = Pt(para.get('gap', el['gap']) * 72)
        first = False
        for text, st in para['runs']:
            r = p.add_run(); r.text = text; f = r.font
            f.name = st.get('font', para.get('font', el['font']))
            f.size = Pt(st.get('size', para.get('size', el['size'])))
            f.bold = st.get('bold', para.get('bold', el['bold']))
            f.italic = st.get('italic', para.get('italic', el['italic']))
            f.color.rgb = rgb(st.get('color', para.get('color', el['color'])))
            if st.get('underline'):
                f.underline = True
            if st.get('link'):
                r.hyperlink.address = st['link']
            rPr = r._r.get_or_add_rPr()
            spc = st.get('spc', para.get('spc', el['spc']))
            if spc:
                rPr.set('spc', str(int(spc)))
            for tag in ('a:ea', 'a:cs'):
                e = rPr.find(qn(tag))
                if e is None:
                    e = etree.SubElement(rPr, qn(tag))
                e.set('typeface', f.name)


def add_el(slide, el):
    sh = slide.shapes; t = el['t']
    if t in ('rect', 'ellipse'):
        kind = MSO_SHAPE.OVAL if t == 'ellipse' else (MSO_SHAPE.ROUNDED_RECTANGLE if el.get('r') else MSO_SHAPE.RECTANGLE)
        s = sh.add_shape(kind, Inches(el['x']), Inches(el['y']), Inches(el['w']), Inches(el['h']))
        strip_style(s)
        if t == 'rect' and el.get('r'):
            s.adjustments[0] = min(0.5, el['r'] / min(el['w'], el['h']))
        if el['fill']:
            s.fill.solid(); s.fill.fore_color.rgb = rgb(el['fill'])
        else:
            s.fill.background()
        if el['line']:
            s.line.color.rgb = rgb(el['line']); s.line.width = Pt(el['lw'])
            if el.get('dash'):
                s.line.dash_style = 4
        else:
            s.line.fill.background()
    elif t == 'text':
        s = sh.add_textbox(Inches(el['x']), Inches(el['y']), Inches(el['w']), Inches(el['h']))
        fill_tf(s.text_frame, el)
    elif t == 'line':
        s = sh.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(el['x1']), Inches(el['y1']), Inches(el['x2']), Inches(el['y2']))
        strip_style(s)
        s.line.color.rgb = rgb(el['color']); s.line.width = Pt(el['lw'])
        ln = s.line._get_or_add_ln()
        if el['dash']:
            etree.SubElement(ln, qn('a:prstDash')).set('val', 'dash')
        if el['head'] in ('end', 'both'):
            te = etree.SubElement(ln, qn('a:tailEnd')); te.set('type', 'triangle'); te.set('w', 'med'); te.set('len', 'med')
        if el['head'] == 'both':
            he = etree.SubElement(ln, qn('a:headEnd')); he.set('type', 'triangle'); he.set('w', 'med'); he.set('len', 'med')
    elif t == 'arrow':
        kind = {'r': MSO_SHAPE.RIGHT_ARROW, 'l': MSO_SHAPE.LEFT_ARROW, 'd': MSO_SHAPE.DOWN_ARROW, 'u': MSO_SHAPE.UP_ARROW}[el['d']]
        s = sh.add_shape(kind, Inches(el['x']), Inches(el['y']), Inches(el['w']), Inches(el['h']))
        strip_style(s)
        s.fill.solid(); s.fill.fore_color.rgb = rgb(el['color']); s.line.fill.background()
    elif t == 'img':
        s = sh.add_picture(el['path'], Inches(el['x']), Inches(el['y']), Inches(el['w']), Inches(el['h']))
        s._element.nvPicPr.cNvPr.set('descr', el.get('alt', ''))
    elif t == 'icon':
        s = sh.add_picture(icon_png(el['name'], el['color'], el['sw']), Inches(el['x']), Inches(el['y']), Inches(el['s']), Inches(el['s']))
        s._element.nvPicPr.cNvPr.set('descr', '')
    else:
        raise ValueError(f'unknown element type {t}')
    return s


def style_title(ph, text, cfg, cover=False):
    box = cfg['box']
    ph.left, ph.top, ph.width, ph.height = [Inches(v) for v in box]
    ph.text_frame.text = ''
    fill_tf(ph.text_frame, T(*box, text, size=cfg['size'], font=cfg['font'], bold=True, color=cfg['color'],
                             valign='b' if cover else 't'))


def _template_notes_sps(template):
    """Placeholder shapes (sldImg + empty body) copied from an existing notes slide, if any."""
    z = zipfile.ZipFile(template)
    names = sorted(n for n in z.namelist() if n.startswith('ppt/notesSlides/notesSlide') and n.endswith('.xml'))
    if not names:
        return [etree.fromstring(NOTES_BODY)]
    x = etree.fromstring(z.read(names[0])); out = []
    for sp in x.iter(qn('p:sp')):
        ph = sp.find('.//' + qn('p:ph'))
        if ph is not None and ph.get('type') in ('sldImg', 'body'):
            sp = copy.deepcopy(sp)
            if ph.get('type') == 'body':
                txb = sp.find(qn('p:txBody'))
                for p_ in txb.findall(qn('a:p')):
                    txb.remove(p_)
                etree.SubElement(txb, qn('a:p'))
            out.append(sp)
    return out or [etree.fromstring(NOTES_BODY)]


def build(slides, out, template, theme=None):
    th = resolve_theme(theme)
    prs = Presentation(template)
    lst = prs.slides._sldIdLst
    for sldId in list(lst):
        prs.part.drop_rel(sldId.rId); lst.remove(sldId)
    for rId, rel in list(prs.part.rels.items()):
        if rel.reltype.endswith('/slide'):
            prs.part.drop_rel(rId)
    layouts = [l for l in prs.slide_layouts if l.name == th['layout']]
    layout = layouts[0] if layouts else prs.slide_layouts[0]
    note_sps = _template_notes_sps(template)
    for n, sd in enumerate(slides, 1):
        slide = prs.slides.add_slide(layout)
        title_ph = None
        for ph in list(slide.placeholders):
            if ph.placeholder_format.type == 1 and title_ph is None:
                title_ph = ph
            else:
                ph._element.getparent().remove(ph._element)
        cover = sd['kind'] == 'title'
        bgc = th['title_bg'] if cover else th['content_bg']
        if bgc:
            fill = (f'<a:schemeClr val="{bgc}"/>' if not all(c in '0123456789ABCDEFabcdef' for c in bgc) or len(bgc) != 6
                    else f'<a:srgbClr val="{bgc}"/>')
            bg = etree.fromstring(f'<p:bg {NS}><p:bgPr><a:solidFill>{fill}</a:solidFill><a:effectLst/></p:bgPr></p:bg>')
            slide._element.find(qn('p:cSld')).insert(0, bg)
        if cover and th['title_hide_master']:
            slide._element.set('showMasterSp', '0')
        tcfg = dict(th['cover_title'] if cover else th['title'])
        if not cover:
            tcfg['box'] = title_box(sd, th)
        if title_ph is None:  # layout without a title placeholder: draw a text box
            add_el(slide, T(*tcfg['box'], sd['title'], size=tcfg['size'], font=tcfg['font'], bold=True, color=tcfg['color']))
        else:
            style_title(title_ph, sd['title'], tcfg, cover)
        for el in sd['els']:
            add_el(slide, el)
        for el in chrome(sd, n, th):
            add_el(slide, el)
        ns = slide.notes_slide
        if ns.notes_placeholder is None:
            for sp in note_sps:
                ns.shapes._spTree.append(copy.deepcopy(sp))
        ns.notes_text_frame.text = sd.get('notes', '')
    prs.save(out)
    return out
