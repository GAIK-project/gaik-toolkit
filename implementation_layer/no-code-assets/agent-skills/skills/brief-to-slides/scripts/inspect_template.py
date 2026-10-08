#!/usr/bin/env python3
"""Extract the visual rules of a user-supplied .pptx/.potx so a THEME can be written for it.

  python inspect_template.py template.pptx [--slides 6] [--render OUTDIR]

Prints: slide size, layouts, master/layout recurring shapes (logos, bands, placeholders),
theme colours and fonts, the most used fonts/sizes/colours/fills/lines on slides, title and
kicker positions, citation-like small text, and notes conventions. With --render it also
writes PNGs + a contact sheet (via render_preview.sh) to look at.
"""
import argparse, collections, os, re, subprocess, zipfile
from pptx import Presentation
from pptx.util import Emu

IN = 914400


def r2(v):
    return round((v or 0) / IN, 2)


def main():
    a = argparse.ArgumentParser(); a.add_argument('path'); a.add_argument('--slides', type=int, default=6); a.add_argument('--render')
    o = a.parse_args()
    p = Presentation(o.path)
    print(f'size: {r2(p.slide_width)} x {r2(p.slide_height)} in (ratio {p.slide_width / p.slide_height:.3f})')
    z = zipfile.ZipFile(o.path)
    for n in sorted(x for x in z.namelist() if re.match(r'ppt/theme/theme\d+\.xml', x)):
        t = z.read(n).decode('utf8', 'ignore')
        cols = re.findall(r'<a:(dk1|lt1|dk2|lt2|accent\d|hlink)>.*?(?:val|lastClr)="(\w+)"', t)
        fonts = re.findall(r'<a:(major|minor)Font><a:latin typeface="([^"]*)"', t)
        print(f'{n}: colours {dict(cols)} fonts {dict(fonts)}')
    for m in p.slide_masters:
        print('\nMASTER shapes:')
        for sh in m.shapes:
            print(f'  {sh.shape_type} {sh.name!r} at ({r2(sh.left)},{r2(sh.top)}) {r2(sh.width)}x{r2(sh.height)}')
        for l in m.slide_layouts:
            used = sum(1 for s in p.slides if s.slide_layout == l)
            pics = [sh.name for sh in l.shapes if sh.shape_type == 13]
            phs = [(str(ph.placeholder_format.type).split('.')[-1].split(' ')[0], r2(ph.left), r2(ph.top), r2(ph.width), r2(ph.height)) for ph in l.placeholders]
            print(f'  LAYOUT {l.name!r} used_by={used} pictures={pics} placeholders={phs}')
    fonts, sizes, colors, fills, lines = (collections.Counter() for _ in range(5))
    for s in p.slides:
        x = s.part.blob.decode('utf8', 'ignore') if hasattr(s.part, 'blob') else ''
        for sh in s.shapes:
            if sh.has_text_frame:
                for para in sh.text_frame.paragraphs:
                    for r in para.runs:
                        if r.font.name: fonts[r.font.name] += 1
                        if r.font.size: sizes[r.font.size.pt] += 1
                        try:
                            if r.font.color and r.font.color.type is not None and r.font.color.rgb: colors[str(r.font.color.rgb)] += 1
                        except Exception:
                            pass
    for n in z.namelist():
        if re.match(r'ppt/slides/slide\d+\.xml', n):
            x = z.read(n).decode('utf8', 'ignore')
            for sp in re.findall(r'<p:spPr>(.*?)</p:spPr>', x, re.S):
                m_ = re.search(r'</a:(?:prstGeom|custGeom)>\s*<a:solidFill><a:srgbClr val="(\w+)"', sp)
                if m_: fills[m_.group(1)] += 1
                m_ = re.search(r'<a:ln[^>]*>\s*<a:solidFill><a:srgbClr val="(\w+)"', sp)
                if m_: lines[m_.group(1)] += 1
    print('\nfonts', fonts.most_common(8)); print('sizes', sorted(sizes.most_common(10)))
    print('text colours', colors.most_common(10)); print('fills', fills.most_common(12)); print('lines', lines.most_common(8))
    print(f'\nFIRST {o.slides} SLIDES (shapes with text, pictures):')
    for i, s in enumerate(p.slides, 1):
        if i > o.slides: break
        print(f'--- slide {i} layout={s.slide_layout.name!r}')
        for sh in s.shapes:
            t = sh.text_frame.text.replace('\n', ' | ')[:80] if sh.has_text_frame else ''
            f = ''
            if sh.has_text_frame and sh.text_frame.paragraphs and sh.text_frame.paragraphs[0].runs:
                r = sh.text_frame.paragraphs[0].runs[0].font
                f = f'{r.name} {r.size.pt if r.size else ""} {"B" if r.bold else ""}'
            if t or sh.shape_type == 13:
                print(f'  ({r2(sh.left)},{r2(sh.top)}) {r2(sh.width)}x{r2(sh.height)} {sh.shape_type} [{f}] {t!r}')
        if s.has_notes_slide and s.notes_slide.notes_text_frame:
            print('  NOTES:', s.notes_slide.notes_text_frame.text[:160].replace('\n', ' / '))
    if o.render:
        subprocess.run(['bash', os.path.join(os.path.dirname(os.path.abspath(__file__)), 'render_preview.sh'), o.path, o.render, '60'])


if __name__ == '__main__':
    main()
