#!/usr/bin/env python3
"""Show which slides differ between two versions of a deck (text, notes, shapes, images).

  python diff_decks.py old.pptx new.pptx

Use after a targeted change to confirm that only the intended slides changed.
Slides are matched by position; added/removed slides are reported at the end.
"""
import hashlib
import sys
from pptx import Presentation


def fingerprint(slide):
    parts = []
    for sh in slide.shapes:
        geo = (sh.shape_type, sh.left, sh.top, sh.width, sh.height)
        txt = sh.text_frame.text if sh.has_text_frame else ''
        if txt.strip().isdigit():          # slide numbers shift when slides are added/removed
            continue
        img = ''
        if sh.shape_type == 13:
            try:
                img = hashlib.md5(sh.image.blob).hexdigest()
            except Exception:
                pass
        parts.append(repr((geo, txt, img)))
    notes = slide.notes_slide.notes_text_frame.text if slide.has_notes_slide and slide.notes_slide.notes_text_frame else ''
    title = next((sh.text_frame.text for sh in slide.shapes if sh.is_placeholder and sh.placeholder_format.type == 1), '')
    return hashlib.md5('\n'.join(parts).encode()).hexdigest(), hashlib.md5(notes.encode()).hexdigest(), title


def main(a, b):
    A = [fingerprint(s) for s in Presentation(a).slides]
    Bv = [fingerprint(s) for s in Presentation(b).slides]
    changed = 0
    for i in range(min(len(A), len(Bv))):
        body, notes = A[i][0] != Bv[i][0], A[i][1] != Bv[i][1]
        if body or notes:
            changed += 1
            what = ' + '.join(x for x, f in (('slide content', body), ('notes', notes)) if f)
            print(f'slide {i + 1}: {what} changed  ({Bv[i][2][:60]!r})')
    if len(Bv) > len(A):
        for i in range(len(A), len(Bv)):
            print(f'slide {i + 1}: added  ({Bv[i][2][:60]!r})')
    elif len(A) > len(Bv):
        print(f'{len(A) - len(Bv)} slide(s) removed at the end (check order if slides were removed in the middle)')
    print(f'{changed} slide(s) changed; {len(A)} -> {len(Bv)} slides')


if __name__ == '__main__':
    main(*sys.argv[1:3])
