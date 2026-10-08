#!/usr/bin/env python3
"""Quality gates for a built deck. Exit code 1 if any check fails.

  python check_deck.py --pptx deck.pptx [--minutes 60] [--preview preview.html]
                       [--forbid "Old title,Old company"] [--min-notes 600]

Checks: speaker notes on every slide (length), timings parsed from notes '[N min' sum to
--minutes, no hidden slides, no forbidden (leftover template) text, no label words on slides
(mock-up, illustrative, fictional, synthetic data), no unfilled [Placeholders], and - with --preview -
text overflow / mid-word wrapping measured in Chromium with the real fonts.
"""
import argparse, asyncio, re, sys
from pptx import Presentation


LABELS = re.compile(r'mock-?up|illustrative|fictional|synthetic data|\(synthetic\)', re.I)
PLACEHOLDER = re.compile(r'\[[A-Z€$£_][^\]\d§]{1,40}\]')  # [Name], [Affiliation], [€__]; not citations like [Policy v2 §3]


def pptx_checks(path, minutes, forbid, min_notes, allow_labels=False):
    issues = []
    prs = Presentation(path)
    total = 0
    for i, s in enumerate(prs.slides, 1):
        if s._element.get('show') == '0':
            issues.append(f'slide {i}: hidden slide')
        notes = s.notes_slide.notes_text_frame.text if s.has_notes_slide and s.notes_slide.notes_text_frame else ''
        m = re.match(r'\s*\[(\d+(?:\.\d+)?)\s*min', notes)
        mins = float(m.group(1)) if m else 0
        total += mins
        need = min_notes if mins > 0 else 200
        if len(notes) < need:
            issues.append(f'slide {i}: notes too short ({len(notes)} < {need} chars)')
        slide_text = ' '.join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame)
        if not allow_labels:
            for m_ in LABELS.finditer(slide_text):
                issues.append(f'slide {i}: label word "{m_.group(0)}" on slide (keep such disclosures in the notes)')
        for m_ in PLACEHOLDER.finditer(slide_text):
            issues.append(f'slide {i}: unfilled placeholder {m_.group(0)}')
        text = slide_text + ' ' + notes
        for f in forbid:
            if f and f.lower() in text.lower():
                issues.append(f'slide {i}: leftover text "{f}"')
    if minutes is not None and abs(total - minutes) > 1e-6:
        issues.append(f'timings sum to {total:g} min, expected {minutes:g}')
    print(f'{len(prs.slides)} slides, timed total {total:g} min')
    return issues


async def overflow(preview):
    from playwright.async_api import async_playwright
    async with async_playwright() as p:
        b = await p.chromium.launch()
        pg = await b.new_page(viewport={'width': 1920, 'height': 1080})
        await pg.goto('file://' + __import__('os').path.abspath(preview)); await pg.wait_for_timeout(800)
        res = await pg.evaluate('''()=>{const out=[];document.querySelectorAll('section').forEach((sec,i)=>{
          sec.querySelectorAll(':scope > div').forEach(d=>{const ps=[...d.children].filter(c=>/^(P|H1|H2)$/.test(c.tagName)); if(!ps.length) return;
            const r=d.getBoundingClientRect(); let tot=(ps.length-1)*(parseFloat(d.style.gap)||0); ps.forEach(p=>tot+=p.getBoundingClientRect().height);
            const wide=ps.some(p=>p.scrollWidth>p.clientWidth+2);
            if(tot>r.height+3||wide) out.push(`slide ${i+1} (${sec.id}): ${wide?'word wider than box':'overflow '+Math.round(tot-r.height)+'px'} :: ${d.innerText.slice(0,50).replace(/\\n/g,' ')}`);});
          sec.querySelectorAll(':scope > div').forEach(d=>{const r=d.getBoundingClientRect(), s=sec.getBoundingClientRect();
            if(r.right>s.right+1||r.bottom>s.bottom+1) out.push(`slide ${i+1} (${sec.id}): element off canvas`);});});return out;}''')
        await b.close()
        return res


def main():
    a = argparse.ArgumentParser()
    a.add_argument('--pptx', required=True); a.add_argument('--minutes', type=float)
    a.add_argument('--preview'); a.add_argument('--forbid', default=''); a.add_argument('--min-notes', type=int, default=600)
    a.add_argument('--allow-labels', action='store_true', help='skip the on-slide label-word check (only if the user wants such labels)')
    o = a.parse_args()
    issues = pptx_checks(o.pptx, o.minutes, [f.strip() for f in o.forbid.split(',')], o.min_notes, o.allow_labels)
    if o.preview:
        issues += asyncio.run(overflow(o.preview))
    for x in issues:
        print('FAIL', x)
    print('OK' if not issues else f'{len(issues)} issue(s)')
    sys.exit(1 if issues else 0)


if __name__ == '__main__':
    main()
