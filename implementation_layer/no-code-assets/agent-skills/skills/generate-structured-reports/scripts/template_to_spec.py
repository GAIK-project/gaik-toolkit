#!/usr/bin/env python3
"""Turn a report template into a draft outline of sections for review.

  python template_to_spec.py TEMPLATE [--out outline.json]

TEMPLATE is a Word (.docx), Markdown (.md) or text (.txt) file; a text-based PDF is read but
rarely has headings. The output is a draft for the user to review, not an executable spec: it
has no sources, and every mapping it is unsure of is marked. The agent merges the approved
outline with the real inputs into report_spec.json and checks it with run_report.py preflight.

Nothing here calls a model.
"""

from __future__ import annotations

import argparse
import json
import os
import re
import sys
import unicodedata
from pathlib import Path

ID_PATTERN = re.compile(r"^[A-Za-z0-9_-]+$")
ATX = re.compile(r"^(#{1,6})\s+(.*?)\s*#*\s*$")
SETEXT = re.compile(r"^(=+|-+)\s*$")
FENCE = re.compile(r"^\s*(```|~~~)")
BULLET = re.compile(r"^\s*(?:[-*+]|\d+[.)])\s+(.*\S)\s*$")
PLACEHOLDER = re.compile(r"\[[^\]]*\]|<[^>]*>|\{[^}]*\}|_{3,}|\.{4,}")
CHECKLIST_INTRO = re.compile(
    r"\b(include|includes|including|cover|covers|must|should|required|list|describe|state|"
    r"record|checklist|sisällytä|kuvaa|mainitse|ilmoita|tulee|pitää)\b",
    re.IGNORECASE,
)
DERIVED_TITLES = re.compile(
    r"^(executive\s+summary|summary|conclusions?|recommendations?|overall\s+assessment|"
    r"findings\s+and\s+recommendations|yhteenveto|johtopäätökset|suositukset|"
    r"toimenpide-?ehdotukset|sammanfattning|slutsatser|rekommendationer)\b",
    re.IGNORECASE,
)
TRANSLITERATE = {"ß": "ss", "æ": "ae", "Æ": "ae", "ø": "o", "Ø": "o", "œ": "oe", "Œ": "oe"}


def slug(title: str) -> str:
    """An ASCII section id: 'Yhteenveto ja jatkotoimet' -> 'yhteenveto_ja_jatkotoimet'."""
    text = "".join(TRANSLITERATE.get(ch, ch) for ch in title)
    text = unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode("ascii")
    text = re.sub(r"[^A-Za-z0-9]+", "_", text.lower()).strip("_")
    return re.sub(r"_+", "_", text)


def docx_markdown(path: Path) -> str:
    """A Word template as Markdown that keeps its headings and its bullet lists.

    gaik's DocxParser drops list markers, so a checklist would read as plain text; a template's
    bullets are what become required items, so they are read here with python-docx instead.
    """
    import docx

    lines: list[str] = []
    for paragraph in docx.Document(str(path)).paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style = (paragraph.style.name if paragraph.style else "") or ""
        level = re.match(r"Heading (\d)", style)
        if style == "Title":
            lines.append(f"# {text}")
        elif level:
            lines.append(f"{'#' * int(level.group(1))} {text}")
        elif (
            style.startswith("List")
            or paragraph._p.pPr is not None
            and (paragraph._p.pPr.numPr is not None)
        ):
            lines.append(f"- {text}")
        else:
            lines.append(text)
        lines.append("")
    return "\n".join(lines)


def to_markdown(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix in (".md", ".txt"):
        return path.read_text(encoding="utf-8-sig")
    if suffix == ".docx":
        return docx_markdown(path)
    if suffix == ".pdf":
        from gaik.software_components.source_normalizer import SourceNormalizer

        # A local parser; no config and no model call.
        return SourceNormalizer().to_markdown(path)
    raise SystemExit(f"{path.name}: a template must be .docx, .md, .txt or a text-based .pdf")


def headings(markdown: str) -> list[tuple[int, int, str]]:
    """(line index, level, title) of every heading outside code fences."""
    lines = markdown.splitlines()
    found, in_fence = [], False
    for i, line in enumerate(lines):
        if FENCE.match(line):
            in_fence = not in_fence
            continue
        if in_fence:
            continue
        match = ATX.match(line)
        if match and match.group(2).strip():
            found.append((i, len(match.group(1)), match.group(2).strip()))
            continue
        if (
            i + 1 < len(lines)
            and SETEXT.match(lines[i + 1])
            and line.strip()
            and not BULLET.match(line)
            and not line.lstrip().startswith(("|", ">"))
        ):
            found.append((i, 1 if lines[i + 1].lstrip().startswith("=") else 2, line.strip()))
    return found


def read_body(lines: list[str]) -> tuple[str, list[dict], list[str], list[str]]:
    """Split a section body into guidance, candidate required items, placeholders, subheadings."""
    guidance, items, placeholders, subheadings = [], [], [], []
    intro = ""
    run: list[str] = []

    def close_run() -> None:
        if not run:
            return
        short = all(len(b.split()) <= 10 and not b.rstrip().endswith(".") for b in run)
        no_facts = not any(re.search(r"\d", b) for b in run)
        explicit = bool(CHECKLIST_INTRO.search(intro))
        kind = "checklist" if explicit or (short and no_facts) else "unclear"
        items.extend({"text": b, "kind": kind} for b in run)
        run.clear()

    in_fence = False
    for line in lines:
        if FENCE.match(line):
            in_fence = not in_fence
            continue
        if in_fence or SETEXT.match(line):
            continue
        match = ATX.match(line)
        if match:
            close_run()
            subheadings.append(match.group(2).strip())
            guidance.append(f"{match.group(2).strip()}:")  # keeps the text under it labelled
            continue
        bullet = BULLET.match(line)
        if bullet:
            run.append(bullet.group(1))
            continue
        close_run()
        text = line.strip()
        if not text:
            continue
        found = PLACEHOLDER.findall(text)
        if found:
            placeholders.extend(found)
        guidance.append(text)
        intro = text
    close_run()
    return " ".join(guidance), items, placeholders, subheadings


def outline(markdown: str, name: str) -> dict:
    lines = markdown.splitlines()
    found = headings(markdown)
    diagnostics: list[dict] = []
    if not found:
        return {
            "template": name,
            "title": None,
            "sections": [],
            "diagnostics": [
                {
                    "severity": "error",
                    "message": "No headings found. Use a template with headings, or propose an "
                    "outline as for a report without a template",
                }
            ],
        }

    title = None
    top = min(level for _, level, _ in found)
    tops = [h for h in found if h[1] == top]
    if len(tops) == 1 and len(found) > 1 and found[0][1] == top:
        title = tops[0][2]
        found = found[1:]
        diagnostics.append(
            {"severity": "info", "message": f"Used {title!r} as the report title, not a section"}
        )
    if not found:
        diagnostics.append({"severity": "error", "message": "The template has only a title"})
        return {"template": name, "title": title, "sections": [], "diagnostics": diagnostics}

    level = min(lv for _, lv, _ in found)
    starts = [(i, lv, t) for i, lv, t in found if lv == level]
    if len({t for _, lv, t in found if lv == level}) > 1 and any(
        b[1] - a[1] > 1 for a, b in zip(found, found[1:], strict=False)
    ):
        diagnostics.append(
            {
                "severity": "warning",
                "message": "Heading levels skip (for example level 2 to level 4); check that the "
                "sections were split where intended",
            }
        )

    sections, used_ids, seen_titles = [], set(), {}
    for n, (start, _, heading) in enumerate(starts):
        end = starts[n + 1][0] if n + 1 < len(starts) else len(lines)
        guidance, items, placeholders, subheadings = read_body(lines[start + 1 : end])
        base = slug(heading)
        notes = []
        if not base:
            base = f"section_{n + 1}"
            notes.append("The title gives no ASCII id; review the generated id")
        # Reserve every id handed out, so 'Findings', 'Findings', 'Findings 2' cannot give two
        # 'findings_2': the same rule as run_report.py set-sections.
        section_id, k = base, 2
        while section_id in used_ids:
            section_id, k = f"{base}_{k}", k + 1
        if section_id != base:
            notes.append(f"The id {base!r} was taken; renamed to {section_id!r}")
        used_ids.add(section_id)
        assert ID_PATTERN.match(section_id)
        key = heading.casefold()
        if key in seen_titles:
            diagnostics.append(
                {"severity": "warning", "message": f"Two sections are titled {heading!r}"}
            )
        seen_titles[key] = True
        if any(i["kind"] == "unclear" for i in items):
            notes.append(
                "Some bullets may be example content rather than required items; confirm them"
            )
        if placeholders:
            notes.append("The guidance has placeholders; treat them as instructions, not facts")
        if subheadings:
            notes.append("Subheadings are kept inside this section, not split into sections")
        sections.append(
            {
                "id": section_id,
                "title": heading,
                "instructions": guidance,
                "candidate_required_items": items,
                "placeholders": placeholders,
                "subheadings": subheadings,
                "suggest_derived": bool(DERIVED_TITLES.match(heading)),
                "suggested_depends_on": [],
                "notes": notes,
            }
        )

    evidence = [s["id"] for s in sections if not s["suggest_derived"]]
    for s in sections:
        if s["suggest_derived"]:
            if evidence:
                s["suggested_depends_on"] = evidence
                s["notes"].append(
                    "The title suggests a section written from the others; confirm, and keep it "
                    "evidence-based if it needs its own source material"
                )
            else:
                s["suggest_derived"] = False
    if not evidence:
        diagnostics.append(
            {
                "severity": "warning",
                "message": "Every section looks like a summary; at least one must be written "
                "from the sources",
            }
        )
    return {
        "template": name,
        "title": title,
        "structure": draft_structure(sections),
        "sections": sections,
        "diagnostics": diagnostics,
    }


def draft_structure(sections: list[dict]) -> list[dict]:
    """The sections in the form the user reviews and `run_report.py set-sections` accepts.

    Only checklist bullets become required items; the unsure ones are listed to confirm. A
    section with no guidance text gets an empty prompt, which the agent drafts from the
    template and the normalized sources before showing it.
    """
    titles = {s["id"]: s["title"] for s in sections}
    draft = []
    for s in sections:
        entry = {
            "title": s["title"],
            "prompt": s["instructions"],
            "required_items": [
                i["text"] for i in s["candidate_required_items"] if i["kind"] == "checklist"
            ],
            "summary_of": [titles[d] for d in s["suggested_depends_on"]],
        }
        unsure = [i["text"] for i in s["candidate_required_items"] if i["kind"] == "unclear"]
        if unsure:
            entry["to_confirm"] = unsure
        draft.append(entry)
    return draft


def main(argv: list[str] | None = None) -> int:
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(encoding="utf-8")
        except (AttributeError, ValueError):
            pass
    parser = argparse.ArgumentParser(description="Draft report sections from a template.")
    parser.add_argument("template", type=Path)
    parser.add_argument("--out", type=Path, help="write the outline here instead of stdout")
    args = parser.parse_args(argv)
    if not args.template.is_file():
        print(f"{args.template}: no such file", file=sys.stderr)
        return 2
    if args.out and args.out.exists() and os.path.samefile(args.out, args.template):
        print(f"{args.out}: --out is the template itself; choose another file", file=sys.stderr)
        return 2
    result = outline(to_markdown(args.template), args.template.name)
    text = json.dumps(result, indent=2, ensure_ascii=False)
    if args.out:
        args.out.write_text(text + "\n", encoding="utf-8")
    else:
        print(text)
    return 2 if any(d["severity"] == "error" for d in result["diagnostics"]) else 0


if __name__ == "__main__":
    sys.exit(main())
