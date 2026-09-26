"""Offline tests for SourceNormalizer and the parser options it relies on.

Fixture files are built in ``tmp_path``. The Transcriber and VisionParser are replaced
by fakes, so no test calls an API.
"""

from __future__ import annotations

from pathlib import Path
from types import SimpleNamespace

import pymupdf
import pytest
from docx import Document
from gaik.software_components.llm.base import UsageCounter
from gaik.software_components.parsers.docx_parser import DocxParser
from gaik.software_components.parsers.pymypdf import PyMuPDFParser
from gaik.software_components.source_normalizer import (
    NormalizedSource,
    NormalizedSources,
    SourceNormalizer,
)
from gaik.software_components.source_normalizer.models import source_id
from openpyxl import Workbook

CONFIG = {"provider": "azure", "model": "gpt-test", "api_key": "fake"}


def _pdf(path: Path, pages: list[str]) -> Path:
    doc = pymupdf.open()
    for text in pages:
        page = doc.new_page()
        if text:
            page.insert_text((72, 72), text)
    doc.save(path)
    doc.close()
    return path


def _docx(path: Path) -> Path:
    doc = Document()
    doc.add_heading("Inspection report", level=0)
    doc.add_heading("Roof", level=1)
    doc.add_paragraph("Tiles are cracked.")
    doc.add_paragraph("")
    table = doc.add_table(rows=2, cols=2)
    for (row, col), text in {(0, 0): "Item", (0, 1): "Cost", (1, 0): "Gutter | east"}.items():
        table.cell(row, col).text = text
    table.cell(1, 1).text = "450"
    doc.add_heading("Walls", level=2)
    doc.save(path)
    return path


def _xlsx(path: Path) -> Path:
    workbook = Workbook()
    workbook.active.title = "Log"
    workbook.active.append(["Item", "Cost"])
    workbook.active.append(["Roof", 450])
    workbook.save(path)
    return path


def _text(path: Path, text: str) -> Path:
    path.write_text(text, encoding="utf-8")
    return path


def _fake(calls: list, text: str) -> type:
    """A stand-in for Transcriber and VisionParser that records its constructor args."""

    class Fake:
        def __init__(self, *args, **kwargs):
            calls.append((args, kwargs))
            self.usage = UsageCounter()

        def transcribe(self, path):
            assert Path(calls[-1][1]["output_dir"]).is_dir()
            return SimpleNamespace(raw_transcript=text, usage={"audio_seconds": 12})

        def convert_image(self, path):
            self.usage.add({"prompt_tokens": 4})
            return text

    return Fake


# ---------------------------------------------------------------------------
# Parser options
# ---------------------------------------------------------------------------


def test_pymupdf_page_markers(tmp_path):
    path = str(_pdf(tmp_path / "a.pdf", ["First page", "Second page"]))
    parser = PyMuPDFParser()

    assert parser.parse_pdf(path, page_markers=True) == (
        "[Page 1]\nFirst page\n\n\n[Page 2]\nSecond page"
    )
    assert parser.parse_pdf(path) == "First page\n\n\nSecond page"
    assert parser.parse_document(path, page_markers=True)["text_content"].startswith("[Page 1]")


def test_docx_keep_structure_walks_body_in_order(tmp_path):
    path = str(_docx(tmp_path / "a.docx"))
    parser = DocxParser()

    assert parser.parse_docx(path, keep_structure=True) == (
        "# Inspection report\n\n# Roof\n\nTiles are cracked.\n\n"
        "| Item | Cost |\n| --- | --- |\n| Gutter \\| east | 450 |\n\n## Walls"
    )
    # Default output is unchanged: paragraphs first, then table rows.
    assert parser.parse_docx(path) == (
        "Inspection report\nRoof\nTiles are cracked.\nWalls\nItem | Cost\nGutter | east | 450"
    )


# ---------------------------------------------------------------------------
# SourceNormalizer
# ---------------------------------------------------------------------------


def test_normalize_dispatches_local_types_in_input_order(tmp_path):
    txt = _text(tmp_path / "notes.txt", "Site visit notes.")
    csv = _text(tmp_path / "costs.csv", "Item,Cost\nRoof,450\n")
    xlsx = _xlsx(tmp_path / "log.xlsx")
    docx = _docx(tmp_path / "report.docx")
    pdf = _pdf(tmp_path / "policy.pdf", ["Policy text"])
    messages: list[str] = []

    sources = SourceNormalizer().normalize(
        {"primary": [txt, str(csv)], "secondary": [xlsx, docx, pdf]},
        progress_callback=messages.append,
    )

    assert [(s.id, s.file, s.source_class, s.source_type, s.tool) for s in sources.sources] == [
        ("01_notes", "notes.txt", "primary", "text", "text"),
        ("02_costs", "costs.csv", "primary", "spreadsheet", "SpreadsheetParser"),
        ("03_log", "log.xlsx", "secondary", "spreadsheet", "SpreadsheetParser"),
        ("04_report", "report.docx", "secondary", "docx", "DocxParser"),
        ("05_policy", "policy.pdf", "secondary", "pdf", "PyMuPDFParser"),
    ]
    assert messages == [
        f"Normalizing {name} ({i}/5)"
        for i, name in enumerate(
            ["notes.txt", "costs.csv", "log.xlsx", "report.docx", "policy.pdf"], start=1
        )
    ]
    assert sources.get("notes.txt").text == "Site visit notes."
    assert sources.get("costs.csv").text.startswith("| Row | Item | Cost |")
    assert sources.get("log.xlsx").text.startswith("## Sheet: Log")
    assert sources.get("report.docx").text.startswith("# Inspection report")
    assert sources.get("policy.pdf").text == "[Page 1]\nPolicy text"


def test_plain_list_has_no_source_class(tmp_path):
    sources = SourceNormalizer().normalize([_text(tmp_path / "a.md", "# A")])

    assert sources.sources[0].source_class is None
    assert sources.sources[0].id == "01_a"


@pytest.mark.parametrize(
    ("index", "name", "expected"),
    [
        (1, "notes.txt", "01_notes"),
        (1, "floor_plan (1).png", "01_floor_plan_1"),
        (
            3,
            "Condition assessment – 12 Example Road, Sampleton.docx",
            "03_Condition_assessment_12_Example_Road_Sampleton",
        ),
        (2, "kuntoarvio_ä.pdf", "02_kuntoarvio"),
        (4, "ääni.mp3", "04_ni"),
        (5, "ääö.mp3", "05"),
        (1, 'a"b\\c:d?.txt', "01_a_b_c_d"),
    ],
)
def test_source_id_is_a_portable_slug(index, name, expected):
    assert source_id(index, name) == expected
    NormalizedSource(id=expected, file=name, source_class=None, source_type="text", tool="t", text="")


def test_any_file_name_keeps_its_name_and_gets_a_slug_id(tmp_path):
    sources = SourceNormalizer().normalize([_text(tmp_path / "notes (1).txt", "Notes.")])

    [source] = sources.sources
    assert (source.id, source.file) == ("01_notes_1", "notes (1).txt")
    assert (tmp_path / "out" / "01_notes_1.md") == sources.save(tmp_path / "out")["01_notes_1"]
    assert NormalizedSources.load(tmp_path / "out").get("notes (1).txt").text == "Notes."


def test_to_markdown_uses_the_same_dispatch(tmp_path):
    path = _docx(tmp_path / "sample.docx")
    assert SourceNormalizer().to_markdown(path).startswith("# Inspection report\n\n# Roof")


@pytest.mark.parametrize(
    ("name", "content", "error", "match"),
    [
        ("deck.pptx", "x", ValueError, r"deck\.pptx: unsupported file type \.pptx"),
        ("old.doc", "x", ValueError, r"old\.doc: unsupported file type \.doc"),
        ("blank.txt", " \n", ValueError, r"blank\.txt: no text extracted"),
        ("talk.mp3", "x", ValueError, r"talk\.mp3: audio files need an LLM config"),
        ("roof.png", "x", ValueError, r"roof\.png: image files need an LLM config"),
    ],
)
def test_bad_files_raise(tmp_path, name, content, error, match):
    with pytest.raises(error, match=match):
        SourceNormalizer().normalize([_text(tmp_path / name, content)])


def test_pdf_without_text_layer_raises(tmp_path):
    path = _pdf(tmp_path / "scan.pdf", ["", ""])
    with pytest.raises(ValueError, match=r"scan\.pdf: PDF has no text layer"):
        SourceNormalizer().normalize([path])


def test_directory_and_missing_path_raise(tmp_path):
    for path in (tmp_path, tmp_path / "missing.txt"):
        with pytest.raises(FileNotFoundError, match="not an existing file"):
            SourceNormalizer().normalize([path])


def test_duplicate_file_names_raise(tmp_path):
    (tmp_path / "a").mkdir()
    (tmp_path / "b").mkdir()
    first = _text(tmp_path / "a" / "notes.txt", "one")
    second = _text(tmp_path / "b" / "notes.txt", "two")
    with pytest.raises(ValueError, match=r"Duplicate source file names.*notes\.txt"):
        SourceNormalizer().normalize({"primary": [first], "secondary": [second]})


def test_all_files_are_checked_before_converting(tmp_path):
    messages: list[str] = []
    paths = [_text(tmp_path / "notes.txt", "ok"), _text(tmp_path / "deck.pptx", "x")]
    with pytest.raises(ValueError, match="unsupported file type"):
        SourceNormalizer().normalize(paths, progress_callback=messages.append)
    assert messages == []


def test_parser_error_names_the_file(tmp_path):
    path = tmp_path / "broken.xlsx"
    path.write_bytes(b"not a zip file")
    with pytest.raises(Exception) as info:
        SourceNormalizer().normalize([path])
    assert f"While normalizing {path}" in info.value.__notes__


def test_audio_uses_transcriber_raw_transcript(tmp_path, monkeypatch):
    calls: list = []
    monkeypatch.setattr(
        "gaik.software_components.transcriber.transcriber.Transcriber",
        _fake(calls, "Hello from the site."),
    )
    path = _text(tmp_path / "visit.m4a", "audio")

    normalizer = SourceNormalizer(CONFIG, transcription_model="gpt-4o-transcribe", language="fi")
    sources = normalizer.normalize({"primary": [path]})
    source = sources.sources[0]
    assert sources.usage == {"audio_seconds": 12}

    assert (source.source_type, source.tool, source.text) == (
        "audio",
        "Transcriber",
        "Hello from the site.",
    )
    args, kwargs = calls[0]
    assert args == (CONFIG,)
    assert kwargs["transcription_model"] == "gpt-4o-transcribe"
    assert kwargs["language"] == "fi"
    assert "enhanced_transcript" not in kwargs


def test_failed_transcript_segment_raises(tmp_path, monkeypatch):
    text = "Part one.\n[Transcription failed for segment 2]"
    monkeypatch.setattr(
        "gaik.software_components.transcriber.transcriber.Transcriber", _fake([], text)
    )
    path = _text(tmp_path / "visit.mp3", "audio")
    with pytest.raises(RuntimeError, match=r"visit\.mp3: transcription failed"):
        SourceNormalizer(CONFIG).normalize([path])


def test_image_uses_vision_parser_with_model_override(tmp_path, monkeypatch):
    calls: list = []
    monkeypatch.setattr(
        "gaik.software_components.parsers.vision.VisionParser", _fake(calls, "A cracked tile.")
    )
    path = _text(tmp_path / "roof.jpg", "image")

    normalizer = SourceNormalizer(CONFIG, vision_model="vision-x")
    sources = normalizer.normalize([path])
    source = sources.sources[0]
    assert sources.usage == {"prompt_tokens": 4}
    assert normalizer.normalize([path]).usage == {"prompt_tokens": 4}  # per call
    assert normalizer.usage.snapshot() == {"prompt_tokens": 8}  # running total
    SourceNormalizer(CONFIG).normalize([path])

    assert (source.source_type, source.tool, source.text) == (
        "image",
        "VisionParser",
        "A cracked tile.",
    )
    assert calls[0][0] == ({**CONFIG, "model": "vision-x"},)
    assert calls[2][0] == (CONFIG,)


def test_save_load_round_trip(tmp_path):
    sources = SourceNormalizer().normalize(
        {"primary": [_text(tmp_path / "notes.txt", "Ääni ja kuva.")], "secondary": []}
    )
    sources.save(tmp_path / "out")

    assert NormalizedSources.load(tmp_path / "out") == sources


@pytest.mark.parametrize("bad_id", ["../x", "a/b", r"a\b", ".hidden", ".."])
def test_source_id_cannot_leave_the_folder(bad_id):
    with pytest.raises(ValueError, match="id"):
        NormalizedSource(
            id=bad_id, file="a.txt", source_class=None, source_type="text", tool="t", text="x"
        )
